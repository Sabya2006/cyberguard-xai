import os
from pptx import Presentation

pptx_path = "CYBERGUARD_XAI_Presentation.pptx"

if not os.path.exists(pptx_path):
    print("ERROR: PPTX file does not exist!")
    exit(1)

prs = Presentation(pptx_path)

print(f"Total Slides Count: {len(prs.slides)}")
assert len(prs.slides) == 18, f"Expected 18 slides, got {len(prs.slides)}"

for idx, slide in enumerate(prs.slides, start=1):
    notes = slide.notes_slide.notes_text_frame.text.strip()
    shape_count = len(slide.shapes)
    print(f"Slide {idx:02d}: {shape_count} shapes, Notes length: {len(notes)} chars")
    assert shape_count > 0, f"Slide {idx} has no shapes!"
    assert len(notes) > 0, f"Slide {idx} has missing speaker notes!"

file_size_kb = os.path.getsize(pptx_path) / 1024.0
print(f"ALL 18 SLIDES VERIFIED SUCCESSFULLY! File size: {file_size_kb:.2f} KB")
